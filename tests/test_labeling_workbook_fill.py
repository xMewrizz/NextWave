"""Tests for filling the labeling workbook template (blue cells only)."""

from __future__ import annotations

import hashlib
import unittest
from datetime import date
from io import BytesIO
from pathlib import Path

import openpyxl

from nextwave.labeling.contracts import EvidenceDirection, EvidenceKind, SourceType, TrustLevel
from nextwave.labeling.queue import (
    LabelingQueue,
    QueuedCandidate,
    QueuedEvidence,
    QueuedNoise,
    RunSearchCoverage,
)
from nextwave.labeling.workbook import (
    count_evidence_rows,
    fill_labeling_workbook,
    read_candidate_slots,
    read_noise_slots,
)

TEMPLATE = Path("templates/labeling_workbook.xlsx")


def empty_queue() -> LabelingQueue:
    return LabelingQueue(
        candidates=(),
        noise=(),
        evidence=(),
        deficits=(),
        overflow=(),
        dropped=(),
        unknown_trust_count=0,
        missing_date_count=0,
        future_evidence_count=0,
    )


def sample_queue() -> LabelingQueue:
    base = empty_queue()
    candidate_slots = read_candidate_slots(TEMPLATE)
    noise_slots = read_noise_slots(TEMPLATE)
    first_candidate = candidate_slots[0]
    first_noise = noise_slots[0]
    from dataclasses import replace

    candidates = (
        QueuedCandidate(
            candidate_id=first_candidate["candidate_id"],
            canonical_name="Speculative Decoding",
            aliases=["speculative-decoding"],
            group_id="alias-group-1",
            source_query="Технологии в ИИ",
            domain=first_candidate["domain"],
            analysis_scope_key="edge-v1",
            run_id="run-1",
        ),
    )
    noise = (
        QueuedNoise(
            noise_id=first_noise["noise_id"],
            planned_noise_type=first_noise["planned_noise_type"],
            source_query="Технологии в ИИ",
            extracted_text="Digital Transformation",
            source_document_url="https://example.org/broad",
            domain="Edge",
            analysis_scope_key="edge-v1",
            duplicate_of_candidate_id=None,
            origin_kind="exclusion",
        ),
    )
    evidence = (
        QueuedEvidence(
            evidence_id="evidence-001",
            candidate_id=first_candidate["candidate_id"],
            direction=EvidenceDirection.SUPPORT,
            kind=EvidenceKind.TECHNICAL_VALIDATION,
            source_type=SourceType.RESEARCH,
            trust_level=TrustLevel.A,
            title="Study",
            url="https://example.org/paper",
            published_at=date(2026, 1, 1),
            organization="",
            origin_id="doi:10.1234/x",
            claim="A prototype was tested.",
            locator="excerpt[0:10]",
            extraction_confidence=None,
        ),
        QueuedEvidence(
            evidence_id="evidence-002",
            candidate_id=first_candidate["candidate_id"],
            direction=EvidenceDirection.COUNTER,
            kind=EvidenceKind.PUBLICITY_WAVE,
            source_type=SourceType.INDUSTRY_MEDIA,
            trust_level=None,
            title="",
            url="https://example.org/news",
            published_at=None,
            organization="",
            origin_id="https://example.org/news",
            claim="Analysts hype the launch.",
            locator="excerpt[0:10]",
            extraction_confidence=None,
        ),
    )
    return replace(
        base,
        candidates=candidates,
        noise=noise,
        evidence=evidence,
        search_coverage=(
            RunSearchCoverage(
                run_id="run-1",
                raw_query="Технологии в ИИ",
                source_classes=("scientific", "industry"),
            ),
        ),
    )


def load_cells(data: bytes, sheet: str, rows: range, columns: str) -> list:
    workbook = openpyxl.load_workbook(BytesIO(data), read_only=True, data_only=True)
    try:
        target = workbook[sheet]
        return [
            [target[f"{column}{row}"].value for column in columns]
            for row in rows
        ]
    finally:
        workbook.close()


class WorkbookSlotTests(unittest.TestCase):
    def test_reads_template_owned_slots(self) -> None:
        candidates = read_candidate_slots(TEMPLATE)
        noise = read_noise_slots(TEMPLATE)

        self.assertEqual(len(candidates), 100)
        self.assertEqual(candidates[0]["candidate_id"], "team-negative-001")
        self.assertEqual(len(noise), 50)
        self.assertTrue(
            all(
                set(row) == {"noise_id", "planned_noise_type", "row_number"}
                for row in noise
            )
        )
        self.assertEqual(count_evidence_rows(TEMPLATE), 400)


class WorkbookFillTests(unittest.TestCase):
    def test_fills_blue_cells_and_preserves_grey_and_yellow(self) -> None:
        before = TEMPLATE.read_bytes()
        filled = fill_labeling_workbook(TEMPLATE, sample_queue())
        after = TEMPLATE.read_bytes()

        self.assertEqual(before, after)
        candidates = load_cells(filled, "Кандидаты", range(5, 105), "ABCDEFGHI")
        first_row = next(row for row in candidates if row[0] == "team-negative-001")
        self.assertEqual(first_row[2], "Speculative Decoding")
        self.assertEqual(first_row[3], "speculative-decoding")
        self.assertEqual(first_row[4], "alias-group-1")
        self.assertEqual(first_row[5], "Технологии в ИИ")
        # Grey cells untouched.
        self.assertEqual(first_row[1], "unclassified")
        self.assertEqual(first_row[6], "Edge")

        noise_rows = load_cells(filled, "Шум", range(5, 55), "ABCDEFGHI")
        first_noise = next(row for row in noise_rows if row[0] == "noise-001")
        self.assertEqual(first_noise[3], "Digital Transformation")
        self.assertEqual(first_noise[4], "https://example.org/broad")
        self.assertEqual(first_noise[5], "Edge")
        self.assertEqual(first_noise[6], "edge-v1")

        evidence_rows = load_cells(filled, "Доказательства", range(5, 7), "ABCDEFGHIJKLMN")
        self.assertEqual(evidence_rows[0][1], "team-negative-001")
        self.assertEqual(evidence_rows[0][2], "support")
        self.assertEqual(evidence_rows[0][5], "A")
        self.assertEqual(evidence_rows[1][2], "counter")
        # Unknown trust stays blank; unknown date stays blank.
        self.assertIsNone(evidence_rows[1][5])
        self.assertIsNone(evidence_rows[1][8])

    def test_decisions_arrive_with_evidence_links_and_search_coverage(self) -> None:
        filled = fill_labeling_workbook(TEMPLATE, sample_queue())
        decisions = load_cells(filled, "Решения", range(5, 205), "ABCDEFGHIJKLMNOP")

        candidate_rows = [row for row in decisions if row[2] == "team-negative-001"]
        self.assertEqual(len(candidate_rows), 1)
        row = candidate_rows[0]
        self.assertEqual(row[11], "evidence-001|evidence-002")
        self.assertEqual(row[12], "Технологии в ИИ")
        self.assertEqual(row[13], "scientific|industry")

        noise_rows = [row for row in decisions if row[2] == "noise-001"]
        self.assertEqual(len(noise_rows), 1)
        self.assertEqual(noise_rows[0][12], "Технологии в ИИ")
        self.assertIsNone(noise_rows[0][11])
        self.assertIsNone(noise_rows[0][13])

    def test_navigation_links_freeze_and_filters_are_present(self) -> None:
        filled = fill_labeling_workbook(TEMPLATE, sample_queue())
        workbook = openpyxl.load_workbook(BytesIO(filled))
        try:
            candidates = workbook["Кандидаты"]
            self.assertIn("Доказательства", candidates["A5"].hyperlink.location)
            evidence = workbook["Доказательства"]
            self.assertIn("Кандидаты", evidence["B5"].hyperlink.location)
            decisions = workbook["Решения"]
            decision_row = next(
                row
                for row in range(5, decisions.max_row + 1)
                if decisions[f"C{row}"].value == "team-negative-001"
            )
            self.assertIn("Кандидаты", decisions[f"C{decision_row}"].hyperlink.location)
            for sheet in (candidates, evidence, decisions, workbook["Шум"]):
                self.assertEqual(sheet.freeze_panes, "A5")
                self.assertTrue(sheet.auto_filter.ref)
        finally:
            workbook.close()

    def test_candidate_without_evidence_links_to_decision_row(self) -> None:
        from dataclasses import replace

        base = empty_queue()
        slots = read_candidate_slots(TEMPLATE)
        first = slots[0]
        queue = replace(
            base,
            candidates=(
                QueuedCandidate(
                    candidate_id=first["candidate_id"],
                    canonical_name="Lonely Tech",
                    aliases=(),
                    group_id="alias-group-9",
                    source_query="q",
                    domain=first["domain"],
                    analysis_scope_key="edge-v1",
                    run_id="run-1",
                ),
            ),
        )
        filled = fill_labeling_workbook(TEMPLATE, queue)
        workbook = openpyxl.load_workbook(BytesIO(filled))
        try:
            candidates = workbook["Кандидаты"]
            location = candidates["A5"].hyperlink.location
            self.assertIn("Решения", location)
        finally:
            workbook.close()

    def test_untouched_sheets_stay_identical(self) -> None:
        filled = fill_labeling_workbook(TEMPLATE, sample_queue())
        for sheet in ("План", "Справочник"):
            original = load_cells(TEMPLATE.read_bytes(), sheet, range(2, 30), "ABCDEF")
            rewritten = load_cells(filled, sheet, range(2, 30), "ABCDEF")
            self.assertEqual(original, rewritten)

    def test_unknown_candidate_id_is_rejected(self) -> None:
        from dataclasses import replace

        queue = replace(
            empty_queue(),
            candidates=(
                QueuedCandidate(
                    candidate_id="team-negative-999",
                    canonical_name="Ghost",
                    aliases=(),
                    group_id="alias-group-9",
                    source_query="q",
                    domain="Edge",
                    analysis_scope_key="edge-v1",
                    run_id="run-1",
                ),
            ),
        )
        with self.assertRaisesRegex(ValueError, "no template row"):
            fill_labeling_workbook(TEMPLATE, queue)

    def test_template_sha256_is_unchanged(self) -> None:
        digest = hashlib.sha256(TEMPLATE.read_bytes()).hexdigest()
        fill_labeling_workbook(TEMPLATE, sample_queue())
        self.assertEqual(hashlib.sha256(TEMPLATE.read_bytes()).hexdigest(), digest)


if __name__ == "__main__":
    unittest.main()
