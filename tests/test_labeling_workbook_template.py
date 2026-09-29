from __future__ import annotations

import unittest
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook

TEMPLATE_PATH = Path(__file__).parents[1] / "templates" / "labeling_workbook.xlsx"


class LabelingWorkbookTemplateTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        cls.workbook = load_workbook(TEMPLATE_PATH, read_only=False, data_only=False)

    @classmethod
    def tearDownClass(cls) -> None:
        cls.workbook.close()

    def test_required_sheets_exist_in_workflow_order(self) -> None:
        self.assertEqual(
            self.workbook.sheetnames,
            ["План", "Кандидаты", "Доказательства", "Решения", "Шум", "Справочник"],
        )

    def test_candidate_slots_match_domain_plan_and_stay_unclassified(self) -> None:
        sheet = self.workbook["Кандидаты"]
        rows = list(sheet.iter_rows(min_row=5, max_row=104, values_only=True))

        self.assertEqual(len(rows), 100)
        self.assertEqual(sheet["B4"].value, "review_pool")
        self.assertEqual(rows[0][0], "team-negative-001")
        self.assertEqual(rows[-1][0], "team-negative-100")
        self.assertEqual(Counter(row[1] for row in rows), {"unclassified": 100})
        review_pool_validation = next(
            item
            for item in sheet.data_validations.dataValidation
            if str(item.sqref) == "B5:B104"
        )
        self.assertEqual(review_pool_validation.formula1, '"unclassified"')
        self.assertEqual(
            Counter(row[6] for row in rows),
            {
                "Edge": 16,
                "Защита ИИ": 16,
                "Индустриальный ИИ": 17,
                "Инфраструктура ИИ": 17,
                "Роботы": 17,
                "Финтех": 17,
            },
        )

    def test_noise_slots_have_ten_rows_per_type(self) -> None:
        sheet = self.workbook["Шум"]
        rows = list(sheet.iter_rows(min_row=5, max_row=54, values_only=True))

        self.assertEqual(len(rows), 50)
        self.assertEqual(rows[0][0], "noise-001")
        self.assertEqual(rows[-1][0], "noise-050")
        self.assertEqual(
            Counter(row[1] for row in rows),
            {
                "broad_concept": 10,
                "irrelevant": 10,
                "not_technology": 10,
                "extraction_error": 10,
                "duplicate": 10,
            },
        )

    def test_noise_domain_cells_are_empty_pipeline_fields(self) -> None:
        sheet = self.workbook["Шум"]
        rows = list(sheet.iter_rows(min_row=5, max_row=54, values_only=False))

        self.assertEqual(len(rows), 50)
        for row in rows:
            domain_cell, scope_cell = row[5], row[6]
            self.assertIsNone(domain_cell.value)
            self.assertIsNone(scope_cell.value)
            self.assertEqual(domain_cell.fill.fgColor.rgb, "FFDDEBF7")
            self.assertEqual(scope_cell.fill.fgColor.rgb, "FFDDEBF7")

    def test_primary_decision_slots_are_prelinked(self) -> None:
        sheet = self.workbook["Решения"]
        candidate_rows = list(sheet.iter_rows(min_row=5, max_row=104, values_only=True))
        noise_rows = list(sheet.iter_rows(min_row=105, max_row=154, values_only=True))

        self.assertTrue(all(row[1] == "model_candidate" for row in candidate_rows))
        self.assertTrue(all(row[3] == "primary" for row in candidate_rows))
        self.assertTrue(all(row[1] == "noise_control" for row in noise_rows))
        self.assertTrue(all(row[3] == "primary" for row in noise_rows))

    def test_evidence_rows_are_pipeline_proposals_for_human_verification(self) -> None:
        sheet = self.workbook["Доказательства"]
        headers = [cell.value for cell in sheet[4]]
        rows = list(sheet.iter_rows(min_row=5, max_row=404, values_only=True))

        self.assertEqual(len(rows), 400)
        self.assertEqual(
            headers[-3:],
            ["extraction_confidence", "verification_status", "review_note"],
        )
        self.assertTrue(all(row[14] == "pending" for row in rows))
        self.assertEqual(sheet["B5"].fill.fgColor.rgb, "FFDDEBF7")
        self.assertEqual(sheet["O5"].fill.fgColor.rgb, "FFFFF2CC")

    def test_workbook_has_validations_tables_and_no_formulas(self) -> None:
        for sheet_name in ("Кандидаты", "Доказательства", "Решения", "Шум"):
            sheet = self.workbook[sheet_name]
            self.assertGreater(len(sheet.data_validations.dataValidation), 0)
            self.assertGreater(len(sheet.tables), 0)
            self.assertIsNotNone(sheet.freeze_panes)
            formulas = [
                cell.coordinate
                for row in sheet.iter_rows()
                for cell in row
                if cell.data_type == "f"
            ]
            self.assertEqual(formulas, [], msg=f"unexpected formulas in {sheet_name}")


if __name__ == "__main__":
    unittest.main()
