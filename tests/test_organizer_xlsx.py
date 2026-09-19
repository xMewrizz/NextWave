from __future__ import annotations

import hashlib
import json
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

from nextwave.datasets import (
    EXPECTED_HEADERS,
    OrganizerWorkbookError,
    build_organizer_dataset,
    read_organizer_workbook,
)


def create_workbook(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Слабые сигналы"
    sheet.cell(1, 2, "Тестовый набор")
    for column, header in enumerate(EXPECTED_HEADERS, start=2):
        sheet.cell(2, column, header)

    domains = (
        "Edge",
        "Защита ИИ",
        "Индустриальный ИИ",
        "Инфраструктура ИИ",
        "Роботы",
        "Финтех",
    )
    for number in range(1, 101):
        row = number + 2
        values = (
            number,
            f"  Технология   {number}  ",
            domains[(number - 1) % len(domains)],
            f"Компания {number}",
            f"Обоснование {number}",
            "Исследование",
            "Рост упоминаний",
            3 + number % 5,
            f"https://example.org/{number}",
        )
        for column, value in enumerate(values, start=2):
            sheet.cell(row, column, value)
    workbook.save(path)


class OrganizerWorkbookTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temporary_directory = tempfile.TemporaryDirectory()
        self.path = Path(self.temporary_directory.name) / "organizer.xlsx"
        create_workbook(self.path)

    def tearDown(self) -> None:
        self.temporary_directory.cleanup()

    def test_valid_workbook_is_parsed_into_separate_records(self) -> None:
        dataset = read_organizer_workbook(self.path)

        self.assertEqual(len(dataset.candidates), 100)
        self.assertEqual(len(dataset.annotations), 100)
        self.assertEqual(dataset.candidates[0].record_id, "organizer-001")
        self.assertEqual(dataset.candidates[-1].record_id, "organizer-100")
        self.assertEqual(dataset.candidates[0].canonical_name, "Технология 1")
        self.assertEqual(dataset.candidates[0].analysis_scope_key, "edge-v1")
        self.assertEqual(dataset.annotations[0].expert_rationale, "Обоснование 1")

    def test_changed_header_is_rejected_with_cell_name(self) -> None:
        workbook = load_workbook(self.path)
        workbook.active.cell(2, 3, "Другое название")
        workbook.save(self.path)

        with self.assertRaisesRegex(OrganizerWorkbookError, "C2"):
            read_organizer_workbook(self.path)

    def test_missing_record_is_rejected(self) -> None:
        workbook = load_workbook(self.path)
        workbook.active.delete_rows(102)
        workbook.save(self.path)

        with self.assertRaisesRegex(OrganizerWorkbookError, "expected data range"):
            read_organizer_workbook(self.path)

    def test_broken_sequence_number_is_rejected(self) -> None:
        workbook = load_workbook(self.path)
        workbook.active.cell(12, 2, 99)
        workbook.save(self.path)

        with self.assertRaisesRegex(OrganizerWorkbookError, "B12"):
            read_organizer_workbook(self.path)

    def test_score_outside_contract_is_rejected(self) -> None:
        workbook = load_workbook(self.path)
        workbook.active.cell(3, 9, 8)
        workbook.save(self.path)

        with self.assertRaisesRegex(OrganizerWorkbookError, "I3"):
            read_organizer_workbook(self.path)

    def test_formula_is_rejected(self) -> None:
        workbook = load_workbook(self.path)
        workbook.active.cell(3, 9, "=3+4")
        workbook.save(self.path)

        with self.assertRaisesRegex(OrganizerWorkbookError, "formula.*I3"):
            read_organizer_workbook(self.path)

    def test_complete_build_contains_verified_manifest(self) -> None:
        paths = build_organizer_dataset(self.path, Path(self.temporary_directory.name) / "output")

        manifest = json.loads(paths.manifest.read_text(encoding="utf-8"))
        self.assertEqual(manifest["schema_version"], "organizer-manifest-v1")
        self.assertEqual(manifest["dataset_version"], "organizer-positive-2026-09-15-v1")
        self.assertEqual(manifest["adapter_version"], "organizer-xlsx-v1")
        self.assertEqual(manifest["cutoff_date"], "2026-09-15")
        self.assertEqual(manifest["input_record_count"], 100)
        self.assertEqual(manifest["accepted_record_count"], 100)
        self.assertEqual(manifest["rejected_record_count"], 0)
        self.assertEqual(manifest["validation_errors"], [])
        self.assertEqual(len(manifest["header_mapping"]), 9)

        source_bytes = self.path.read_bytes()
        self.assertEqual(manifest["source"]["size_bytes"], len(source_bytes))
        self.assertEqual(
            manifest["source"]["sha256"], hashlib.sha256(source_bytes).hexdigest()
        )
        expected_outputs = {
            paths.candidates.name: paths.candidates,
            paths.annotations.name: paths.annotations,
        }
        for output in manifest["outputs"]:
            artifact = expected_outputs[output["filename"]]
            content = artifact.read_bytes()
            self.assertEqual(output["size_bytes"], len(content))
            self.assertEqual(output["sha256"], hashlib.sha256(content).hexdigest())


if __name__ == "__main__":
    unittest.main()
