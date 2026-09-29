from __future__ import annotations

import hashlib
import json
import tempfile
import unittest
from collections import Counter
from datetime import date
from pathlib import Path

import openpyxl

from nextwave.labeling.contracts import NegativeCandidateRecord
from nextwave.labeling.enrichment_plan import _bundle_id
from nextwave.labeling.finalize import finalize_labeling_bundle

DOMAINS = (
    ("Edge", "edge-v1", 8),
    ("Защита ИИ", "ai-security-v1", 8),
    ("Индустриальный ИИ", "industrial-ai-v1", 9),
    ("Инфраструктура ИИ", "ai-infrastructure-v1", 9),
    ("Роботы", "robotics-v1", 8),
    ("Финтех", "fintech-v1", 8),
)
NOISE_TYPES = (
    "broad_concept",
    "irrelevant",
    "not_technology",
    "extraction_error",
    "duplicate",
)


def jsonl(records: list[dict]) -> bytes:
    return (
        "\n".join(
            json.dumps(item, ensure_ascii=False, separators=(",", ":"))
            for item in records
        )
        + "\n"
    ).encode()


def digest(data: bytes) -> dict:
    return {"size_bytes": len(data), "sha256": hashlib.sha256(data).hexdigest()}


def make_bundle(root: Path) -> tuple[Path, Path, Path]:
    bundle = root / "bundle"
    bundle.mkdir()
    candidates = []
    labels: dict[str, str] = {}
    number = 1
    for domain, scope, mature_count in DOMAINS:
        total = 16 if domain in {"Edge", "Защита ИИ"} else 17
        for index in range(total):
            candidate_id = f"team-negative-{number:03d}"
            candidates.append(
                {
                    "schema_version": "negative-candidate-v1",
                    "candidate_id": candidate_id,
                    "canonical_name": f"Technology {number}",
                    "aliases": [],
                    "group_id": f"group-{number:03d}",
                    "source_query": f"Query {domain}",
                    "domain": domain,
                    "analysis_scope_key": scope,
                    "cutoff_date": "2026-09-15",
                }
            )
            labels[candidate_id] = "mature" if index < mature_count else "marketing_hype"
            number += 1
    noise = []
    for number in range(1, 51):
        domain, scope, _ = DOMAINS[(number - 1) % len(DOMAINS)]
        noise.append(
            {
                "schema_version": "noise-control-v1",
                "noise_id": f"noise-{number:03d}",
                "source_query": f"Query {domain}",
                "extracted_text": f"Noise {number}",
                "source_document_url": f"https://example.org/noise/{number}",
                "domain": domain,
                "analysis_scope_key": scope,
                "cutoff_date": "2026-09-15",
                "duplicate_of_candidate_id": None,
            }
        )
    candidate_bytes = jsonl(candidates)
    noise_bytes = jsonl(noise)
    (bundle / "negative_candidates.jsonl").write_bytes(candidate_bytes)
    (bundle / "noise_controls.jsonl").write_bytes(noise_bytes)
    manifest = {
        "schema_version": "labeling-export-manifest-v1",
        "rubric_version": "labeling-v1",
        "cutoff_date": "2026-09-15",
        "candidate_selection": {"sha256": "candidate-selection"},
        "noise_selection": {"sha256": "noise-selection"},
        "deficits": [],
        "outputs": {
            "negative_candidates.jsonl": digest(candidate_bytes),
            "noise_controls.jsonl": digest(noise_bytes),
        },
    }
    (bundle / "manifest.json").write_text(
        json.dumps(manifest, ensure_ascii=False), encoding="utf-8"
    )
    workbook_path = root / "review.xlsx"
    make_workbook(workbook_path, candidates, noise, labels)
    enrichment = make_enrichment_result(root, candidates)
    return bundle, workbook_path, enrichment


def make_enrichment_result(root: Path, candidates: list[dict]) -> Path:
    result = root / "enrichment"
    result.mkdir()
    coverage = []
    for candidate in candidates:
        for connector, source_class, planned in (
            ("openalex", "scientific", 2),
            ("mediacloud", "industry", 1),
        ):
            coverage.append(
                {
                    "candidate_id": candidate["candidate_id"],
                    "connector": connector,
                    "failed_request_ids": [],
                    "failed_requests": 0,
                    "incompleteness_reasons": [],
                    "parse_issue_count": 0,
                    "planned_requests": planned,
                    "returned_documents": 0,
                    "returned_records": 0,
                    "reused_requests": 0,
                    "search_id": f"search-{candidate['candidate_id']}-{source_class}",
                    "source_class": source_class,
                    "status": "complete",
                    "successful_requests": planned,
                }
            )
    coverage_bytes = jsonl(coverage)
    (result / "coverage.jsonl").write_bytes(coverage_bytes)
    manifest = {
        "schema_version": "labeling-enrichment-result-v2",
        "bundle_id": _bundle_id(
            [
                NegativeCandidateRecord(
                    candidate_id=item["candidate_id"],
                    canonical_name=item["canonical_name"],
                    aliases=tuple(item["aliases"]),
                    group_id=item["group_id"],
                    source_query=item["source_query"],
                    domain=item["domain"],
                    analysis_scope_key=item["analysis_scope_key"],
                    cutoff_date=date.fromisoformat(item["cutoff_date"]),
                )
                for item in candidates
            ]
        ),
        "totals": {"candidates": len(candidates)},
        "outputs": {"coverage.jsonl": digest(coverage_bytes)},
    }
    (result / "manifest.json").write_text(
        json.dumps(manifest, ensure_ascii=False), encoding="utf-8"
    )
    return result


def make_workbook(
    path: Path,
    candidates: list[dict],
    noise: list[dict],
    labels: dict[str, str],
) -> None:
    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    candidate_sheet = workbook.create_sheet("Кандидаты")
    evidence_sheet = workbook.create_sheet("Доказательства")
    noise_sheet = workbook.create_sheet("Шум")
    decisions = workbook.create_sheet("Решения")
    for sheet, headers in (
        (
            candidate_sheet,
            (
                "candidate_id",
                "review_pool",
                "canonical_name",
                "aliases",
                "group_id",
                "source_query",
                "domain",
                "analysis_scope_key",
                "cutoff_date",
            ),
        ),
        (
            evidence_sheet,
            (
                "evidence_id",
                "candidate_id",
                "direction",
                "kind",
                "source_type",
                "trust_level",
                "title",
                "url",
                "published_at",
                "organization",
                "origin_id",
                "claim",
                "locator",
                "confidence",
                "verification_status",
                "review_note",
            ),
        ),
        (
            noise_sheet,
            (
                "noise_id",
                "planned_noise_type",
                "source_query",
                "extracted_text",
                "source_document_url",
                "domain",
                "analysis_scope_key",
                "cutoff_date",
                "duplicate_of_candidate_id",
            ),
        ),
        (
            decisions,
            (
                "decision_id",
                "subject_type",
                "subject_id",
                "review_round",
                "status",
                "model_label",
                "noise_type",
                "rationale",
                "reviewer_id",
                "annotated_at",
                "cutoff_date",
                "evidence_ids",
                "search_queries",
                "search_source_classes",
                "searched_at",
                "search_notes",
            ),
        ),
    ):
        for column, value in enumerate(headers, 1):
            sheet.cell(4, column, value)

    evidence_for: dict[str, list[str]] = {}
    evidence_row = 5
    for index, candidate in enumerate(candidates, 1):
        candidate_sheet.append(
            (
                candidate["candidate_id"],
                "unclassified",
                candidate["canonical_name"],
                "",
                candidate["group_id"],
                candidate["source_query"],
                candidate["domain"],
                candidate["analysis_scope_key"],
                date(2026, 9, 15),
            )
        )
        evidence_ids = []
        count = 1 if labels[candidate["candidate_id"]] == "mature" else 3
        for item_index in range(count):
            evidence_id = f"evidence-{evidence_row - 4:03d}"
            evidence_ids.append(evidence_id)
            is_mature = labels[candidate["candidate_id"]] == "mature"
            origin_suffix = "a" if item_index < 2 else "b"
            values = (
                evidence_id,
                candidate["candidate_id"],
                "support",
                "serial_deployment" if is_mature else "publicity_wave",
                "official_technical" if is_mature else "industry_media",
                "A" if is_mature else "D",
                f"Evidence {evidence_id}",
                f"https://example.org/evidence/{evidence_id}",
                date(2026, 7, 1 + item_index),
                "Example Org",
                f"origin-{index:03d}-{origin_suffix}",
                "A dated and verifiable claim.",
                "section-1",
                None,
                "verified",
                None,
            )
            for column, value in enumerate(values, 1):
                evidence_sheet.cell(evidence_row, column, value)
            evidence_row += 1
        evidence_for[candidate["candidate_id"]] = evidence_ids

    for index, item in enumerate(noise, 1):
        noise_type = NOISE_TYPES[(index - 1) // 10]
        noise_sheet.append(
            (
                item["noise_id"],
                noise_type,
                item["source_query"],
                item["extracted_text"],
                item["source_document_url"],
                item["domain"],
                item["analysis_scope_key"],
                date(2026, 9, 15),
                None,
            )
        )

    decision_number = 1
    secondary_model = {"mature": 0, "marketing_hype": 0}
    for candidate in candidates:
        candidate_id = candidate["candidate_id"]
        label = labels[candidate_id]
        append_decision(
            decisions,
            decision_number,
            "model_candidate",
            candidate_id,
            "primary",
            label,
            None,
            "reviewer-primary",
            evidence_for[candidate_id],
        )
        decision_number += 1
        if secondary_model[label] < 10:
            append_decision(
                decisions,
                decision_number,
                "model_candidate",
                candidate_id,
                "secondary",
                label,
                None,
                "reviewer-secondary",
                evidence_for[candidate_id],
            )
            decision_number += 1
            secondary_model[label] += 1
    secondary_noise = Counter()
    for index, item in enumerate(noise, 1):
        noise_type = NOISE_TYPES[(index - 1) // 10]
        append_decision(
            decisions,
            decision_number,
            "noise_control",
            item["noise_id"],
            "primary",
            None,
            noise_type,
            "reviewer-primary",
            [],
        )
        decision_number += 1
        if secondary_noise[noise_type] < 2:
            append_decision(
                decisions,
                decision_number,
                "noise_control",
                item["noise_id"],
                "secondary",
                None,
                noise_type,
                "reviewer-secondary",
                [],
            )
            decision_number += 1
            secondary_noise[noise_type] += 1
    workbook.save(path)
    workbook.close()


def append_decision(
    sheet,
    number: int,
    subject_type: str,
    subject_id: str,
    review_round: str,
    model_label: str | None,
    noise_type: str | None,
    reviewer: str,
    evidence_ids: list[str],
) -> None:
    is_model = subject_type == "model_candidate"
    values = (
        f"decision-{number:03d}",
        subject_type,
        subject_id,
        review_round,
        "reviewed",
        model_label,
        noise_type,
        "The evidence satisfies the current rubric.",
        reviewer,
        date(2026, 9, 27),
        date(2026, 9, 15),
        "|".join(evidence_ids),
        "technology|deployment" if is_model else None,
        "scientific|industry" if is_model else None,
        date(2026, 9, 27) if is_model else None,
        "scientific=complete;industry=complete" if is_model else None,
    )
    for column, value in enumerate(values, 1):
        sheet.cell(sheet.max_row + (1 if column == 1 else 0), column, value)


class LabelingFinalizeTests(unittest.TestCase):
    def test_publishes_balanced_reviewed_corpus(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            bundle, workbook, enrichment = make_bundle(root)
            paths = finalize_labeling_bundle(
                bundle_dir=bundle,
                workbook_path=workbook,
                enrichment_result_dir=enrichment,
                output_dir=root / "output",
            )
            manifest = json.loads(paths.manifest.read_text(encoding="utf-8"))
            self.assertEqual(manifest["counts"]["mature"], 50)
            self.assertEqual(manifest["counts"]["marketing_hype"], 50)
            self.assertEqual(manifest["counts"]["noise_controls"], 50)
            self.assertEqual(manifest["secondary_review"]["mature"]["reviewed"], 10)
            rows = [
                json.loads(line)
                for line in paths.candidates.read_text(encoding="utf-8").splitlines()
            ]
            self.assertTrue(all(item["target"] == 0 for item in rows))

    def test_rejects_missing_secondary_review_without_output(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            bundle, workbook, enrichment = make_bundle(root)
            book = openpyxl.load_workbook(workbook)
            sheet = book["Решения"]
            for row in range(5, sheet.max_row + 1):
                if sheet[f"D{row}"].value == "secondary":
                    sheet[f"D{row}"].value = "primary"
            book.save(workbook)
            book.close()
            output = root / "output"
            with self.assertRaisesRegex(ValueError, "invalid decision history"):
                finalize_labeling_bundle(
                    bundle_dir=bundle,
                    workbook_path=workbook,
                    enrichment_result_dir=enrichment,
                    output_dir=output,
                )
            self.assertFalse(output.exists())

    def test_does_not_treat_workbook_note_as_coverage_source(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            bundle, workbook, enrichment = make_bundle(root)
            book = openpyxl.load_workbook(workbook)
            sheet = book["Решения"]
            row = next(
                row
                for row in range(5, sheet.max_row + 1)
                if sheet[f"F{row}"].value == "marketing_hype"
            )
            sheet[f"P{row}"].value = "scientific=partial;industry=complete"
            book.save(workbook)
            book.close()
            paths = finalize_labeling_bundle(
                bundle_dir=bundle,
                workbook_path=workbook,
                enrichment_result_dir=enrichment,
                output_dir=root / "output",
            )
            self.assertTrue(paths.manifest.is_file())

    def test_rejects_workbook_from_another_bundle(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            bundle, workbook, enrichment = make_bundle(root)
            book = openpyxl.load_workbook(workbook)
            book["Кандидаты"]["C5"] = "Changed identity"
            book.save(workbook)
            book.close()
            with self.assertRaisesRegex(ValueError, "differs from bundle"):
                finalize_labeling_bundle(
                    bundle_dir=bundle,
                    workbook_path=workbook,
                    enrichment_result_dir=enrichment,
                    output_dir=root / "output",
                )

    def test_rejects_candidate_level_partial_coverage(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            bundle, workbook, enrichment = make_bundle(root)
            rows = [
                json.loads(line)
                for line in (enrichment / "coverage.jsonl")
                .read_text(encoding="utf-8")
                .splitlines()
            ]
            rows[0]["status"] = "partial"
            rows[0]["successful_requests"] = 1
            rows[0]["failed_requests"] = 1
            rows[0]["failed_request_ids"] = ["request-failed"]
            payload = jsonl(rows)
            (enrichment / "coverage.jsonl").write_bytes(payload)
            manifest = json.loads(
                (enrichment / "manifest.json").read_text(encoding="utf-8")
            )
            manifest["outputs"]["coverage.jsonl"] = digest(payload)
            (enrichment / "manifest.json").write_text(
                json.dumps(manifest), encoding="utf-8"
            )
            with self.assertRaisesRegex(ValueError, "incomplete scientific coverage"):
                finalize_labeling_bundle(
                    bundle_dir=bundle,
                    workbook_path=workbook,
                    enrichment_result_dir=enrichment,
                    output_dir=root / "output",
                )

    def test_rejects_enrichment_from_another_candidate_bundle(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            bundle, workbook, enrichment = make_bundle(root)
            manifest = json.loads(
                (enrichment / "manifest.json").read_text(encoding="utf-8")
            )
            manifest["bundle_id"] = "bundle-0000000000000000"
            (enrichment / "manifest.json").write_text(
                json.dumps(manifest), encoding="utf-8"
            )
            with self.assertRaisesRegex(ValueError, "another candidate bundle"):
                finalize_labeling_bundle(
                    bundle_dir=bundle,
                    workbook_path=workbook,
                    enrichment_result_dir=enrichment,
                    output_dir=root / "output",
                )

    def test_rejects_noise_labels_swapped_between_reviewed_types(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            bundle, workbook, enrichment = make_bundle(root)
            book = openpyxl.load_workbook(workbook)
            sheet = book["Решения"]
            for row in range(5, sheet.max_row + 1):
                if sheet[f"C{row}"].value == "noise-001":
                    sheet[f"G{row}"] = "irrelevant"
                elif sheet[f"C{row}"].value == "noise-011":
                    sheet[f"G{row}"] = "broad_concept"
            book.save(workbook)
            book.close()
            with self.assertRaisesRegex(ValueError, "reviewed selection"):
                finalize_labeling_bundle(
                    bundle_dir=bundle,
                    workbook_path=workbook,
                    enrichment_result_dir=enrichment,
                    output_dir=root / "output",
                )

    def test_rejects_bundle_without_reviewed_noise_selection(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            bundle, workbook, enrichment = make_bundle(root)
            manifest = json.loads(
                (bundle / "manifest.json").read_text(encoding="utf-8")
            )
            manifest["noise_selection"] = None
            (bundle / "manifest.json").write_text(
                json.dumps(manifest), encoding="utf-8"
            )
            with self.assertRaisesRegex(ValueError, "reviewed noise selection"):
                finalize_labeling_bundle(
                    bundle_dir=bundle,
                    workbook_path=workbook,
                    enrichment_result_dir=enrichment,
                    output_dir=root / "output",
                )


if __name__ == "__main__":
    unittest.main()
