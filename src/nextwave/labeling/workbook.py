"""Fill the labeling workbook template from a review queue (blue cells only)."""

from __future__ import annotations

from datetime import date, datetime
from io import BytesIO
from pathlib import Path
from typing import Any
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo

import openpyxl
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.hyperlink import Hyperlink

from .queue import LabelingQueue

# Фиксированная метка сборки ради побайтовой воспроизводимости экспорта.
BUILD_STAMP = datetime(2026, 9, 15)
_ZIP_STAMP = (1980, 1, 1, 0, 0, 0)

CANDIDATES_SHEET = "Кандидаты"
EVIDENCE_SHEET = "Доказательства"
NOISE_SHEET = "Шум"
DECISIONS_SHEET = "Решения"

HEADER_ROW = 4
FIRST_DATA_ROW = 5

LINK_FONT = Font(color="0563C1", underline="single")


def read_candidate_slots(template_path: str | Path) -> list[dict[str, Any]]:
    """Read neutral candidate review rows owned by the template."""

    workbook = openpyxl.load_workbook(template_path, read_only=True, data_only=True)
    try:
        sheet = workbook[CANDIDATES_SHEET]
        slots = []
        for row in sheet.iter_rows(min_row=FIRST_DATA_ROW, values_only=True):
            if not row[0]:
                continue
            if row[1] != "unclassified":
                raise ValueError(
                    f"candidate row {len(slots) + FIRST_DATA_ROW} must use "
                    "the neutral review_pool value 'unclassified'"
                )
            slots.append(
                {
                    "candidate_id": str(row[0]),
                    "domain": str(row[6]),
                    "row_number": len(slots) + FIRST_DATA_ROW,
                }
            )
        return slots
    finally:
        workbook.close()


def read_noise_slots(template_path: str | Path) -> list[dict[str, Any]]:
    """Read template-owned noise rows: ID, planned type and domain per row."""

    workbook = openpyxl.load_workbook(template_path, read_only=True, data_only=True)
    try:
        sheet = workbook[NOISE_SHEET]
        slots = []
        for row in sheet.iter_rows(min_row=FIRST_DATA_ROW, values_only=True):
            if not row[0]:
                continue
            slots.append(
                {
                    "noise_id": str(row[0]),
                    "planned_noise_type": str(row[1]),
                    "domain": str(row[5]),
                    "row_number": len(slots) + FIRST_DATA_ROW,
                }
            )
        return slots
    finally:
        workbook.close()


def count_evidence_rows(template_path: str | Path) -> int:
    """Count template-owned evidence rows available for queue proposals."""

    workbook = openpyxl.load_workbook(template_path, read_only=True, data_only=True)
    try:
        sheet = workbook[EVIDENCE_SHEET]
        total = 0
        for row in sheet.iter_rows(min_row=FIRST_DATA_ROW, values_only=True):
            if not row[0]:
                break
            total += 1
        return total
    finally:
        workbook.close()


def _set_cell(sheet: Any, row: int, column: str, value: Any) -> None:
    if value is None:
        return
    if isinstance(value, str) and not value.strip():
        return
    sheet[f"{column}{row}"] = value


def fill_labeling_workbook(
    template_path: str | Path, queue: LabelingQueue
) -> bytes:
    """Fill a copy of the template with queue data; template file untouched.

    Only blue (pipeline) cells are written. Grey identifiers, review pools,
    domains, cutoff dates and all yellow expert cells stay byte-identical,
    except decision rows, which arrive pre-filled with evidence links and
    machine search coverage (columns L/M/N/P), plus navigation hyperlinks on ID
    cells, freeze panes and filters. The reviewer deletes unused evidence
    links instead of typing identifiers by hand.
    """

    workbook = openpyxl.load_workbook(template_path)
    workbook.properties.created = BUILD_STAMP
    workbook.properties.modified = BUILD_STAMP
    candidates = workbook[CANDIDATES_SHEET]
    candidate_rows: dict[str, int] = {}
    for item in queue.candidates:
        row = _candidate_row(candidates, item.candidate_id)
        if row is None:
            raise ValueError(f"candidate {item.candidate_id} has no template row")
        candidate_rows[item.candidate_id] = row
        _set_cell(candidates, row, "C", item.canonical_name)
        _set_cell(candidates, row, "D", "; ".join(item.aliases))
        _set_cell(candidates, row, "E", item.group_id)
        _set_cell(candidates, row, "F", item.source_query)

    noise = workbook[NOISE_SHEET]
    noise_rows: dict[str, int] = {}
    for item in queue.noise:
        row = _noise_row(noise, item.noise_id)
        if row is None:
            raise ValueError(f"noise {item.noise_id} has no template row")
        noise_rows[item.noise_id] = row
        _set_cell(noise, row, "C", item.source_query)
        _set_cell(noise, row, "D", item.extracted_text)
        _set_cell(noise, row, "E", item.source_document_url)
        _set_cell(noise, row, "I", item.duplicate_of_candidate_id)

    evidence = workbook[EVIDENCE_SHEET]
    evidence_first_row: dict[str, int] = {}
    for index, item in enumerate(queue.evidence):
        row = FIRST_DATA_ROW + index
        cell_id = evidence[f"A{row}"].value
        if not cell_id:
            raise ValueError("queue evidence exceeds template evidence rows")
        evidence_first_row.setdefault(item.candidate_id, row)
        _set_cell(evidence, row, "B", item.candidate_id)
        _set_cell(evidence, row, "C", item.direction.value)
        _set_cell(evidence, row, "D", item.kind.value)
        _set_cell(evidence, row, "E", item.source_type.value)
        _set_cell(
            evidence, row, "F", item.trust_level.value if item.trust_level else None
        )
        _set_cell(evidence, row, "G", item.title)
        _set_cell(evidence, row, "H", item.url)
        if isinstance(item.published_at, date):
            evidence[f"I{row}"] = item.published_at
        _set_cell(evidence, row, "J", item.organization)
        _set_cell(evidence, row, "K", item.origin_id)
        _set_cell(evidence, row, "L", item.claim)
        _set_cell(evidence, row, "M", item.locator)
        if item.extraction_confidence is not None:
            evidence[f"N{row}"] = item.extraction_confidence

    _fill_decision_links(workbook, queue, candidate_rows, noise_rows)
    _apply_navigation(workbook, candidate_rows, noise_rows, evidence_first_row)

    stream = BytesIO()
    workbook.save(stream)
    workbook.close()
    return _normalize_zip(stream.getvalue())


def _normalize_zip(data: bytes) -> bytes:
    """Repack an xlsx with fixed container metadata for byte determinism.

    openpyxl leaves varying bytes in zip headers between processes even for
    identical sheets, which breaks reproducible checksums. Content is untouched;
    only entry timestamps and attributes are pinned.
    """

    source = BytesIO(data)
    target = BytesIO()
    with ZipFile(source) as reader:
        names = reader.namelist()
        payloads = [(name, reader.read(name)) for name in names]
    with ZipFile(target, "w", compression=ZIP_DEFLATED) as writer:
        for name, payload in payloads:
            info = ZipInfo(name, date_time=_ZIP_STAMP)
            info.compress_type = ZIP_DEFLATED
            writer.writestr(info, payload)
    return target.getvalue()


def _decision_rows(sheet: Any, subject_id: str) -> list[int]:
    return [
        row
        for row in range(FIRST_DATA_ROW, sheet.max_row + 1)
        if sheet[f"C{row}"].value == subject_id
    ]


def _fill_decision_links(
    workbook: Any,
    queue: LabelingQueue,
    candidate_rows: dict[str, int],
    noise_rows: dict[str, int],
) -> None:
    """Pre-fill decision rows: evidence links plus machine search coverage.

    The reviewer deletes unused evidence links instead of typing identifiers.
    Search columns arrive filled from the discovery run; the reviewer appends
    manual queries with "|" instead of guessing classes and dates.
    """

    decisions = workbook[DECISIONS_SHEET]
    coverage = {item.run_id: item for item in queue.search_coverage}
    evidence_ids: dict[str, list[str]] = {}
    for item in queue.evidence:
        evidence_ids.setdefault(item.candidate_id, []).append(item.evidence_id)
    candidates = {item.candidate_id: item for item in queue.candidates}
    noise = {item.noise_id: item for item in queue.noise}
    subjects = set(candidates) | set(noise)
    for subject in sorted(subjects):
        item = candidates.get(subject, noise.get(subject))
        search = coverage.get(item.run_id) if subject in candidates else None
        for row in _decision_rows(decisions, subject):
            _set_cell(decisions, row, "L", "|".join(evidence_ids.get(subject, ())))
            _set_cell(decisions, row, "M", item.source_query)
            _set_cell(
                decisions,
                row,
                "N",
                "|".join(search.source_classes) if search else None,
            )
            _set_cell(
                decisions,
                row,
                "P",
                search.notes if search and search.notes else None,
            )


def _link(sheet: Any, row: int, column: str, target_sheet: str, target_row: int) -> None:
    # Explicit object: plain string assignment misparses quoted sheet names
    # into external relationships (location lost).
    cell = sheet[f"{column}{row}"]
    cell.hyperlink = Hyperlink(
        ref=cell.coordinate, location=f"'{target_sheet}'!A{target_row}"
    )
    cell.font = LINK_FONT


def _apply_navigation(
    workbook: Any,
    candidate_rows: dict[str, int],
    noise_rows: dict[str, int],
    evidence_first_row: dict[str, int],
) -> None:
    """Hyperlinks between the three sheets plus freeze panes and filters."""

    candidates = workbook[CANDIDATES_SHEET]
    evidence = workbook[EVIDENCE_SHEET]
    decisions = workbook[DECISIONS_SHEET]
    noise = workbook[NOISE_SHEET]
    decision_rows: dict[str, int] = {}
    for row in range(FIRST_DATA_ROW, decisions.max_row + 1):
        subject = decisions[f"C{row}"].value
        if subject and subject not in decision_rows:
            decision_rows[subject] = row
    for candidate_id, row in candidate_rows.items():
        first = evidence_first_row.get(candidate_id)
        if first is not None:
            _link(candidates, row, "A", EVIDENCE_SHEET, first)
        elif candidate_id in decision_rows:
            _link(candidates, row, "A", DECISIONS_SHEET, decision_rows[candidate_id])
    for candidate_id, first in evidence_first_row.items():
        target = candidate_rows.get(candidate_id)
        if target is not None:
            _link(evidence, first, "B", CANDIDATES_SHEET, target)
    for row in range(FIRST_DATA_ROW, decisions.max_row + 1):
        subject = decisions[f"C{row}"].value
        if not subject:
            continue
        target = candidate_rows.get(subject)
        sheet_name = CANDIDATES_SHEET
        if target is None:
            target = noise_rows.get(subject)
            sheet_name = NOISE_SHEET
        if target is not None:
            _link(decisions, row, "C", sheet_name, target)
    for sheet in (candidates, evidence, decisions, noise):
        sheet.freeze_panes = "A5"
        last_column = get_column_letter(sheet.max_column)
        sheet.auto_filter.ref = f"A{HEADER_ROW}:{last_column}{sheet.max_row}"


def _candidate_row(sheet: Any, candidate_id: str) -> int | None:
    for row in range(FIRST_DATA_ROW, sheet.max_row + 1):
        if sheet[f"A{row}"].value == candidate_id:
            return row
    return None


def _noise_row(sheet: Any, noise_id: str) -> int | None:
    for row in range(FIRST_DATA_ROW, sheet.max_row + 1):
        if sheet[f"A{row}"].value == noise_id:
            return row
    return None
