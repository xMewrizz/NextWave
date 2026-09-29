"""Validate a reviewed labeling workbook and publish the training corpus."""

from __future__ import annotations

import hashlib
import json
import math
from collections import Counter, defaultdict
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Any

import openpyxl

from ..datasets.artifacts import publish_artifact_bundle
from .contracts import (
    LABELING_CUTOFF_DATE,
    RUBRIC_VERSION,
    EvidenceDirection,
    EvidenceKind,
    GateLabelDecision,
    LabelEvidence,
    ModelLabelDecision,
    NegativeCandidateRecord,
    NegativeClass,
    NoiseControlRecord,
    NoiseType,
    ReviewRound,
    ReviewStatus,
    SearchCoverage,
    SearchSourceClass,
    SourceType,
    TrustLevel,
)
from .enrichment_plan import _bundle_id
from .enrichment_run import ENRICHMENT_RESULT_VERSION
from .export import LABELING_EXPORT_MANIFEST_VERSION
from .workbook import (
    CANDIDATES_SHEET,
    DECISIONS_SHEET,
    EVIDENCE_SHEET,
    FIRST_DATA_ROW,
    NOISE_SHEET,
)

LABELING_FINALIZE_VERSION = "labeling-finalize-v1"
LABELED_NEGATIVES_FILENAME = "labeled_negative_candidates.jsonl"
LABEL_DECISIONS_FILENAME = "label_decisions.jsonl"
VERIFIED_NOISE_FILENAME = "verified_noise_controls.jsonl"
MANIFEST_FILENAME = "manifest.json"

_MODEL_QUOTAS = {
    "Edge": {"mature": 8, "marketing_hype": 8},
    "Защита ИИ": {"mature": 8, "marketing_hype": 8},
    "Индустриальный ИИ": {"mature": 9, "marketing_hype": 8},
    "Инфраструктура ИИ": {"mature": 9, "marketing_hype": 8},
    "Роботы": {"mature": 8, "marketing_hype": 9},
    "Финтех": {"mature": 8, "marketing_hype": 9},
}
_NOISE_QUOTAS = {item.value: 10 for item in NoiseType}


@dataclass(frozen=True, slots=True)
class LabelingFinalizePaths:
    candidates: Path
    decisions: Path
    noise: Path
    manifest: Path


def _digest(data: bytes) -> dict[str, Any]:
    return {"size_bytes": len(data), "sha256": hashlib.sha256(data).hexdigest()}


def _load_json(path: Path) -> dict[str, Any]:
    value = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(value, dict):
        raise ValueError(f"{path.name} must contain a JSON object")
    return value


def _load_jsonl(path: Path) -> list[dict[str, Any]]:
    records: list[dict[str, Any]] = []
    for line_number, line in enumerate(path.read_text(encoding="utf-8").splitlines(), 1):
        value = json.loads(line)
        if not isinstance(value, dict):
            raise ValueError(f"{path.name} line {line_number} must be an object")
        records.append(value)
    if not records:
        raise ValueError(f"{path.name} must not be empty")
    return records


def _verify_input_file(bundle: Path, manifest: dict[str, Any], filename: str) -> Path:
    path = bundle / filename
    if not path.is_file():
        raise FileNotFoundError(f"labeling bundle is missing {filename}")
    expected = (manifest.get("outputs") or {}).get(filename)
    if not isinstance(expected, dict):
        raise ValueError(f"labeling manifest is missing output metadata for {filename}")
    actual = _digest(path.read_bytes())
    if actual != expected:
        raise ValueError(f"{filename} does not match labeling manifest")
    return path


def _validated_candidate_records(
    values: list[dict[str, Any]],
) -> list[NegativeCandidateRecord]:
    records: list[NegativeCandidateRecord] = []
    for index, item in enumerate(values, 1):
        try:
            record = NegativeCandidateRecord(
                candidate_id=item["candidate_id"],
                canonical_name=item["canonical_name"],
                aliases=tuple(item["aliases"]),
                group_id=item["group_id"],
                source_query=item["source_query"],
                domain=item["domain"],
                analysis_scope_key=item["analysis_scope_key"],
                cutoff_date=date.fromisoformat(item["cutoff_date"]),
            )
        except (KeyError, TypeError, ValueError) as error:
            raise ValueError(f"negative candidate line {index} is invalid") from error
        if record.to_dict() != item:
            raise ValueError(f"negative candidate line {index} has unexpected fields")
        records.append(record)
    return records


def _validated_noise_records(values: list[dict[str, Any]]) -> list[NoiseControlRecord]:
    records: list[NoiseControlRecord] = []
    for index, item in enumerate(values, 1):
        try:
            record = NoiseControlRecord(
                noise_id=item["noise_id"],
                source_query=item["source_query"],
                extracted_text=item["extracted_text"],
                source_document_url=item["source_document_url"],
                domain=item["domain"],
                analysis_scope_key=item["analysis_scope_key"],
                cutoff_date=date.fromisoformat(item["cutoff_date"]),
                duplicate_of_candidate_id=item.get("duplicate_of_candidate_id"),
            )
        except (KeyError, TypeError, ValueError) as error:
            raise ValueError(f"noise control line {index} is invalid") from error
        if record.to_dict() != item:
            raise ValueError(f"noise control line {index} has unexpected fields")
        records.append(record)
    return records


def _load_complete_candidate_coverage(
    result_dir: Path,
    candidate_ids: set[str],
    expected_bundle_id: str,
) -> dict[str, dict[str, dict[str, Any]]]:
    manifest_path = result_dir / MANIFEST_FILENAME
    if not manifest_path.is_file():
        raise FileNotFoundError(f"enrichment result is missing {MANIFEST_FILENAME}")
    manifest = _load_json(manifest_path)
    if manifest.get("schema_version") != ENRICHMENT_RESULT_VERSION:
        raise ValueError("enrichment result version does not match current executor")
    if manifest.get("bundle_id") != expected_bundle_id:
        raise ValueError("enrichment result belongs to another candidate bundle")
    expected_candidates = (manifest.get("totals") or {}).get("candidates")
    if type(expected_candidates) is not int or expected_candidates != len(candidate_ids):
        raise ValueError("enrichment candidate count does not match labeling bundle")
    coverage_path = result_dir / "coverage.jsonl"
    expected_digest = (manifest.get("outputs") or {}).get("coverage.jsonl")
    if not coverage_path.is_file() or not isinstance(expected_digest, dict):
        raise ValueError("enrichment result is missing coverage.jsonl metadata")
    if _digest(coverage_path.read_bytes()) != expected_digest:
        raise ValueError("enrichment coverage does not match its manifest")

    by_candidate: dict[str, dict[str, dict[str, Any]]] = defaultdict(dict)
    for item in _load_jsonl(coverage_path):
        candidate_id = item.get("candidate_id")
        source_class = item.get("source_class")
        if candidate_id not in candidate_ids:
            raise ValueError(f"enrichment coverage has unknown candidate {candidate_id!r}")
        if source_class not in {"scientific", "industry"}:
            raise ValueError("enrichment coverage has an unsupported source class")
        if source_class in by_candidate[candidate_id]:
            raise ValueError(
                f"duplicate enrichment coverage for {candidate_id}/{source_class}"
            )
        for field in ("planned_requests", "successful_requests", "failed_requests"):
            if type(item.get(field)) is not int or item[field] < 0:
                raise ValueError(f"enrichment coverage {field} must be a non-negative int")
        if (
            item.get("status") != "complete"
            or item["planned_requests"] == 0
            or item["successful_requests"] != item["planned_requests"]
            or item["failed_requests"] != 0
            or item.get("failed_request_ids") != []
        ):
            raise ValueError(
                f"candidate {candidate_id} has incomplete {source_class} coverage"
            )
        by_candidate[candidate_id][source_class] = item
    if set(by_candidate) != candidate_ids:
        missing = sorted(candidate_ids - set(by_candidate))
        raise ValueError(f"enrichment coverage is missing candidates: {missing[:5]}")
    for candidate_id, classes in by_candidate.items():
        if set(classes) != {"scientific", "industry"}:
            raise ValueError(
                f"candidate {candidate_id} must have scientific and industry coverage"
            )
    return by_candidate


def _render_jsonl(records: list[dict[str, Any]]) -> bytes:
    return (
        "\n".join(
            json.dumps(item, ensure_ascii=False, allow_nan=False, separators=(",", ":"))
            for item in records
        )
        + "\n"
    ).encode("utf-8")


def _text(value: Any, field: str) -> str:
    if not isinstance(value, str) or not value.strip():
        raise ValueError(f"{field} must not be blank")
    return value.strip()


def _optional_text(value: Any) -> str | None:
    if value is None:
        return None
    if not isinstance(value, str):
        raise ValueError("optional text cell must contain text")
    stripped = value.strip()
    return stripped or None


def _date(value: Any, field: str) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        try:
            return date.fromisoformat(value.strip())
        except ValueError as error:
            raise ValueError(f"{field} must use YYYY-MM-DD") from error
    raise ValueError(f"{field} must be a date")


def _split(value: Any, field: str) -> tuple[str, ...]:
    text = _optional_text(value)
    if text is None:
        return ()
    values = tuple(item.strip() for item in text.split("|") if item.strip())
    if len(values) != len(set(values)):
        raise ValueError(f"{field} must not contain duplicates")
    return values


def _load_evidence(
    sheet: Any, candidate_ids: set[str]
) -> dict[str, tuple[str, LabelEvidence]]:
    evidence: dict[str, tuple[str, LabelEvidence]] = {}
    for row, values in enumerate(
        sheet.iter_rows(min_row=FIRST_DATA_ROW, max_col=16, values_only=True),
        FIRST_DATA_ROW,
    ):
        evidence_id = _optional_text(values[0])
        if evidence_id is None:
            continue
        candidate_id = _optional_text(values[1])
        status = _optional_text(values[14])
        if candidate_id is None and status is None:
            continue
        if candidate_id not in candidate_ids:
            raise ValueError(f"evidence row {row} references unknown candidate {candidate_id!r}")
        if status != "verified":
            continue
        if evidence_id in evidence:
            raise ValueError(f"duplicate verified evidence ID {evidence_id}")
        item = LabelEvidence(
            evidence_id=evidence_id,
            direction=EvidenceDirection(_text(values[2], "direction")),
            kind=EvidenceKind(_text(values[3], "kind")),
            source_type=SourceType(_text(values[4], "source_type")),
            trust_level=TrustLevel(_text(values[5], "trust_level")),
            title=_text(values[6], "title"),
            url=_text(values[7], "url"),
            published_at=_date(values[8], "published_at"),
            organization=_text(values[9], "organization"),
            origin_id=_text(values[10], "origin_id"),
            claim=_text(values[11], "claim"),
            locator=_text(values[12], "locator"),
        )
        evidence[evidence_id] = (candidate_id, item)
    return evidence


def _check_workbook_records(
    workbook: Any,
    candidates: dict[str, dict[str, Any]],
    noise: dict[str, dict[str, Any]],
) -> dict[str, NoiseType]:
    candidate_rows: dict[str, tuple[Any, ...]] = {}
    sheet = workbook[CANDIDATES_SHEET]
    for row in sheet.iter_rows(min_row=FIRST_DATA_ROW, values_only=True):
        if row[0]:
            candidate_rows[str(row[0])] = row
    if set(candidate_rows) != set(candidates):
        raise ValueError("reviewed workbook candidates do not match labeling bundle")
    for candidate_id, item in candidates.items():
        row = candidate_rows[candidate_id]
        expected = (
            item["canonical_name"],
            "; ".join(item.get("aliases") or []),
            item["group_id"],
            item["source_query"],
            item["domain"],
            item["analysis_scope_key"],
            item["cutoff_date"],
        )
        actual = (
            row[2],
            row[3] or "",
            row[4],
            row[5],
            row[6],
            row[7],
            _date(row[8], "candidate cutoff").isoformat(),
        )
        if actual != expected:
            raise ValueError(f"workbook candidate {candidate_id} differs from bundle")

    noise_rows: dict[str, tuple[Any, ...]] = {}
    sheet = workbook[NOISE_SHEET]
    for row in sheet.iter_rows(min_row=FIRST_DATA_ROW, values_only=True):
        if row[0]:
            noise_rows[str(row[0])] = row
    if set(noise_rows) != set(noise):
        raise ValueError("reviewed workbook noise rows do not match labeling bundle")
    planned_types: dict[str, NoiseType] = {}
    for noise_id, item in noise.items():
        row = noise_rows[noise_id]
        try:
            planned_type = NoiseType(_text(row[1], "planned noise type"))
        except ValueError as error:
            raise ValueError(f"workbook noise {noise_id} has invalid planned type") from error
        expected = (
            item["source_query"],
            item["extracted_text"],
            item["source_document_url"],
            item["domain"],
            item["analysis_scope_key"],
            item["cutoff_date"],
            item.get("duplicate_of_candidate_id"),
        )
        actual = (
            row[2],
            row[3],
            row[4],
            row[5],
            row[6],
            _date(row[7], "noise cutoff").isoformat(),
            row[8],
        )
        if actual != expected:
            raise ValueError(f"workbook noise {noise_id} differs from bundle")
        planned_types[noise_id] = planned_type
    return planned_types


def _decision_rows(sheet: Any) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for row, raw_values in enumerate(
        sheet.iter_rows(min_row=FIRST_DATA_ROW, max_col=16, values_only=True),
        FIRST_DATA_ROW,
    ):
        values = list(raw_values)
        populated = [value for value in values[1:] if value not in (None, "")]
        if not populated:
            continue
        rows.append(
            {
                "row": row,
                "decision_id": _text(values[0], "decision_id"),
                "subject_type": _text(values[1], "subject_type"),
                "subject_id": _text(values[2], "subject_id"),
                "review_round": _text(values[3], "review_round"),
                "status": _text(values[4], "status"),
                "model_label": _optional_text(values[5]),
                "noise_type": _optional_text(values[6]),
                "rationale": _text(values[7], "rationale"),
                "reviewer_id": _text(values[8], "reviewer_id"),
                "annotated_at": _date(values[9], "annotated_at"),
                "cutoff_date": _date(values[10], "cutoff_date"),
                "evidence_ids": _split(values[11], "evidence_ids"),
                "search_queries": _split(values[12], "search_queries"),
                "search_source_classes": _split(values[13], "search_source_classes"),
                "searched_at": None if values[14] is None else _date(values[14], "searched_at"),
                "search_notes": _optional_text(values[15]),
            }
        )
    return rows


def _model_decision(
    row: dict[str, Any], evidence: dict[str, tuple[str, LabelEvidence]]
) -> ModelLabelDecision:
    if row["subject_type"] != "model_candidate":
        raise ValueError(f"decision row {row['row']} has an invalid subject type")
    if row["noise_type"] is not None:
        raise ValueError(f"model decision row {row['row']} must not set noise_type")
    evidence_items: list[LabelEvidence] = []
    for evidence_id in row["evidence_ids"]:
        stored = evidence.get(evidence_id)
        if stored is None:
            raise ValueError(
                f"decision row {row['row']} references unverified evidence {evidence_id}"
            )
        candidate_id, item = stored
        if candidate_id != row["subject_id"]:
            raise ValueError(
                f"decision row {row['row']} references evidence owned by {candidate_id}"
            )
        evidence_items.append(item)
    coverage = None
    if row["search_queries"] or row["search_source_classes"] or row["searched_at"]:
        if row["searched_at"] is None:
            raise ValueError(f"decision row {row['row']} is missing searched_at")
        coverage = SearchCoverage(
            queries=row["search_queries"],
            source_classes=tuple(
                SearchSourceClass(item) for item in row["search_source_classes"]
            ),
            searched_at=row["searched_at"],
            notes=_text(row["search_notes"], "search_notes"),
        )
    label = NegativeClass(_text(row["model_label"], "model_label"))
    return ModelLabelDecision(
        decision_id=row["decision_id"],
        candidate_id=row["subject_id"],
        review_round=ReviewRound(row["review_round"]),
        status=ReviewStatus(row["status"]),
        label=label,
        rationale=row["rationale"],
        reviewer_id=row["reviewer_id"],
        annotated_at=row["annotated_at"],
        cutoff_date=row["cutoff_date"],
        evidence=tuple(evidence_items),
        search_coverage=coverage,
    )


def _gate_decision(row: dict[str, Any]) -> GateLabelDecision:
    if row["subject_type"] != "noise_control":
        raise ValueError(f"decision row {row['row']} has an invalid subject type")
    if row["model_label"] is not None:
        raise ValueError(f"noise decision row {row['row']} must not set model_label")
    if row["evidence_ids"]:
        raise ValueError(f"noise decision row {row['row']} must not reference evidence")
    return GateLabelDecision(
        decision_id=row["decision_id"],
        noise_id=row["subject_id"],
        review_round=ReviewRound(row["review_round"]),
        status=ReviewStatus(row["status"]),
        noise_type=NoiseType(_text(row["noise_type"], "noise_type")),
        rationale=row["rationale"],
        reviewer_id=row["reviewer_id"],
        annotated_at=row["annotated_at"],
    )


def _resolve_subject(decisions: list[Any], subject_id: str) -> Any:
    primary = [item for item in decisions if item.review_round is ReviewRound.PRIMARY]
    secondary = [item for item in decisions if item.review_round is ReviewRound.SECONDARY]
    adjudication = [item for item in decisions if item.review_round is ReviewRound.ADJUDICATION]
    if len(primary) != 1 or len(secondary) > 1 or len(adjudication) > 1:
        raise ValueError(f"{subject_id} has an invalid decision history")
    if primary[0].status is not ReviewStatus.REVIEWED:
        raise ValueError(f"{subject_id} primary decision must be reviewed")
    if secondary and secondary[0].status is not ReviewStatus.REVIEWED:
        raise ValueError(f"{subject_id} secondary decision must be reviewed")
    if primary[0].reviewer_id in {item.reviewer_id for item in secondary}:
        raise ValueError(f"{subject_id} secondary review must use another reviewer")
    if adjudication and adjudication[0].reviewer_id in {
        primary[0].reviewer_id,
        *(item.reviewer_id for item in secondary),
    }:
        raise ValueError(f"{subject_id} adjudication must use a third reviewer")
    label = getattr(primary[0], "label", getattr(primary[0], "noise_type", None))
    if secondary:
        second_label = getattr(
            secondary[0], "label", getattr(secondary[0], "noise_type", None)
        )
        if second_label != label:
            if len(adjudication) != 1:
                raise ValueError(f"{subject_id} disagreement requires adjudication")
            return adjudication[0]
    if adjudication:
        raise ValueError(f"{subject_id} has adjudication without disagreement")
    return primary[0]


def _validate_secondary_coverage(
    histories: dict[str, list[Any]], final: dict[str, Any], labels: set[str]
) -> dict[str, dict[str, int]]:
    coverage: dict[str, dict[str, int]] = {}
    for label in sorted(labels):
        subject_ids = [
            subject_id
            for subject_id, decision in final.items()
            if str(getattr(decision, "label", getattr(decision, "noise_type", "")))
            == label
        ]
        reviewed = sum(
            any(item.review_round is ReviewRound.SECONDARY for item in histories[item])
            for item in subject_ids
        )
        required = math.ceil(len(subject_ids) * 0.2)
        if reviewed < required:
            raise ValueError(
                f"secondary review for {label} is {reviewed}/{len(subject_ids)}; "
                f"at least {required} is required"
            )
        coverage[label] = {"total": len(subject_ids), "reviewed": reviewed, "required": required}
    return coverage


def finalize_labeling_bundle(
    *,
    bundle_dir: str | Path,
    workbook_path: str | Path,
    enrichment_result_dir: str | Path,
    output_dir: str | Path,
) -> LabelingFinalizePaths:
    """Validate all review decisions and publish a balanced immutable corpus."""

    bundle = Path(bundle_dir)
    workbook_file = Path(workbook_path)
    manifest = _load_json(bundle / "manifest.json")
    if manifest.get("schema_version") != LABELING_EXPORT_MANIFEST_VERSION:
        raise ValueError("input is not a labeling export bundle")
    if manifest.get("rubric_version") != RUBRIC_VERSION:
        raise ValueError("labeling rubric version does not match current rubric")
    if manifest.get("cutoff_date") != LABELING_CUTOFF_DATE.isoformat():
        raise ValueError("labeling cutoff does not match current cutoff")
    if not isinstance(manifest.get("candidate_selection"), dict):
        raise ValueError("final labeling bundle requires reviewed candidate selection")
    if not isinstance(manifest.get("noise_selection"), dict):
        raise ValueError("final labeling bundle requires reviewed noise selection")
    if manifest.get("deficits") != []:
        raise ValueError("final labeling bundle must not contain queue deficits")
    candidates_path = _verify_input_file(bundle, manifest, "negative_candidates.jsonl")
    noise_path = _verify_input_file(bundle, manifest, "noise_controls.jsonl")
    if not workbook_file.is_file():
        raise FileNotFoundError(f"reviewed workbook is missing: {workbook_file}")

    candidates = _load_jsonl(candidates_path)
    noise = _load_jsonl(noise_path)
    candidate_records = _validated_candidate_records(candidates)
    _validated_noise_records(noise)
    candidate_by_id = {item.get("candidate_id"): item for item in candidates}
    noise_by_id = {item.get("noise_id"): item for item in noise}
    if len(candidate_by_id) != len(candidates):
        raise ValueError("negative candidate IDs must be unique")
    if len(noise_by_id) != len(noise):
        raise ValueError("noise IDs must be unique")
    coverage = _load_complete_candidate_coverage(
        Path(enrichment_result_dir),
        set(candidate_by_id),
        _bundle_id(candidate_records),
    )

    workbook = openpyxl.load_workbook(workbook_file, read_only=True, data_only=True)
    try:
        planned_noise = _check_workbook_records(
            workbook, candidate_by_id, noise_by_id
        )
        evidence = _load_evidence(workbook[EVIDENCE_SHEET], set(candidate_by_id))
        rows = _decision_rows(workbook[DECISIONS_SHEET])
    finally:
        workbook.close()

    decision_ids = [row["decision_id"] for row in rows]
    if len(decision_ids) != len(set(decision_ids)):
        raise ValueError("decision IDs must be unique")
    model_histories: dict[str, list[ModelLabelDecision]] = defaultdict(list)
    gate_histories: dict[str, list[GateLabelDecision]] = defaultdict(list)
    all_decisions: list[dict[str, Any]] = []
    for row in rows:
        if row["subject_type"] == "model_candidate":
            decision = _model_decision(row, evidence)
            if decision.candidate_id not in candidate_by_id:
                raise ValueError(f"decision references unknown candidate {decision.candidate_id}")
            model_histories[decision.candidate_id].append(decision)
        elif row["subject_type"] == "noise_control":
            decision = _gate_decision(row)
            if decision.noise_id not in noise_by_id:
                raise ValueError(f"decision references unknown noise {decision.noise_id}")
            gate_histories[decision.noise_id].append(decision)
        else:
            raise ValueError(f"decision row {row['row']} has an invalid subject type")
        all_decisions.append(decision.to_dict())

    if set(model_histories) != set(candidate_by_id):
        missing = sorted(set(candidate_by_id) - set(model_histories))
        raise ValueError(f"model decisions are incomplete: {missing[:5]}")
    if set(gate_histories) != set(noise_by_id):
        missing = sorted(set(noise_by_id) - set(gate_histories))
        raise ValueError(f"noise decisions are incomplete: {missing[:5]}")

    final_model = {
        subject_id: _resolve_subject(history, subject_id)
        for subject_id, history in model_histories.items()
    }
    final_noise = {
        subject_id: _resolve_subject(history, subject_id)
        for subject_id, history in gate_histories.items()
    }
    for noise_id, decision in final_noise.items():
        if decision.noise_type is not planned_noise[noise_id]:
            raise ValueError(
                f"noise decision for {noise_id} does not match reviewed selection"
            )
    by_domain: dict[str, Counter[str]] = defaultdict(Counter)
    for candidate_id, decision in final_model.items():
        by_domain[str(candidate_by_id[candidate_id]["domain"])][decision.label.value] += 1
    if {domain: dict(counts) for domain, counts in by_domain.items()} != _MODEL_QUOTAS:
        raise ValueError("reviewed model decisions do not match the 50/50 domain quotas")
    noise_counts = Counter(item.noise_type.value for item in final_noise.values())
    if dict(noise_counts) != _NOISE_QUOTAS:
        raise ValueError("reviewed noise decisions do not match the five 10-row quotas")

    secondary = {}
    secondary.update(
        _validate_secondary_coverage(
            model_histories, final_model, {item.value for item in NegativeClass}
        )
    )
    secondary.update(
        _validate_secondary_coverage(
            gate_histories, final_noise, {item.value for item in NoiseType}
        )
    )

    labeled = []
    for candidate_id in sorted(candidate_by_id):
        item = dict(candidate_by_id[candidate_id])
        decision = final_model[candidate_id]
        item.update(
            {
                "target": 0,
                "negative_class": decision.label.value,
                "final_decision_id": decision.decision_id,
            }
        )
        labeled.append(item)
    verified_noise = []
    for noise_id in sorted(noise_by_id):
        item = dict(noise_by_id[noise_id])
        decision = final_noise[noise_id]
        item.update(
            {"noise_type": decision.noise_type.value, "final_decision_id": decision.decision_id}
        )
        verified_noise.append(item)
    all_decisions.sort(key=lambda item: item["decision_id"])

    candidates_bytes = _render_jsonl(labeled)
    decisions_bytes = _render_jsonl(all_decisions)
    noise_bytes = _render_jsonl(verified_noise)
    manifest_value = {
        "schema_version": LABELING_FINALIZE_VERSION,
        "rubric_version": RUBRIC_VERSION,
        "cutoff_date": LABELING_CUTOFF_DATE.isoformat(),
        "source_bundle": {
            "schema_version": manifest["schema_version"],
            "manifest_sha256": hashlib.sha256((bundle / "manifest.json").read_bytes()).hexdigest(),
        },
        "enrichment_result": {
            "schema_version": ENRICHMENT_RESULT_VERSION,
            "manifest_sha256": hashlib.sha256(
                (Path(enrichment_result_dir) / MANIFEST_FILENAME).read_bytes()
            ).hexdigest(),
            "complete_candidate_pairs": sum(len(item) for item in coverage.values()),
        },
        "reviewed_workbook": _digest(workbook_file.read_bytes()),
        "counts": {
            "model_candidates": len(labeled),
            "mature": sum(item["negative_class"] == "mature" for item in labeled),
            "marketing_hype": sum(
                item["negative_class"] == "marketing_hype" for item in labeled
            ),
            "noise_controls": len(verified_noise),
            "decisions": len(all_decisions),
            "verified_evidence": len(evidence),
        },
        "model_by_domain": {domain: dict(counts) for domain, counts in sorted(by_domain.items())},
        "noise_by_type": dict(sorted(noise_counts.items())),
        "secondary_review": secondary,
        "outputs": {
            LABELED_NEGATIVES_FILENAME: _digest(candidates_bytes),
            LABEL_DECISIONS_FILENAME: _digest(decisions_bytes),
            VERIFIED_NOISE_FILENAME: _digest(noise_bytes),
        },
    }
    manifest_bytes = (
        json.dumps(manifest_value, ensure_ascii=False, indent=2, sort_keys=True) + "\n"
    ).encode("utf-8")
    paths = publish_artifact_bundle(
        {
            LABELED_NEGATIVES_FILENAME: candidates_bytes,
            LABEL_DECISIONS_FILENAME: decisions_bytes,
            VERIFIED_NOISE_FILENAME: noise_bytes,
            MANIFEST_FILENAME: manifest_bytes,
        },
        output_dir,
    )
    return LabelingFinalizePaths(
        candidates=paths[LABELED_NEGATIVES_FILENAME],
        decisions=paths[LABEL_DECISIONS_FILENAME],
        noise=paths[VERIFIED_NOISE_FILENAME],
        manifest=paths[MANIFEST_FILENAME],
    )
