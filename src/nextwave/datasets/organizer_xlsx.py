"""Strict reader for the organizer-provided positive dataset."""

from __future__ import annotations

from dataclasses import dataclass
from datetime import date
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from .contracts import (
    ORGANIZER_SCOPE_KEYS,
    IdentityStatus,
    OrganizerAnnotationRecord,
    PositiveCandidateRecord,
)

SHEET_NAME = "Слабые сигналы"
HEADER_ROW = 2
DATA_START_ROW = 3
EXPECTED_RECORD_COUNT = 100
DATA_END_ROW = DATA_START_ROW + EXPECTED_RECORD_COUNT - 1
FIRST_COLUMN = 2
LAST_COLUMN = 10
ORGANIZER_CUTOFF_DATE = date(2026, 9, 15)

EXPECTED_HEADERS = (
    "№",
    "Технология (слабый сигнал)",
    "Область",
    "Компании",
    "Почему это слабый сигнал",
    "Стадия развития",
    "Тренд упоминаний",
    "Балл (стадия+тренд)",
    "Источники",
)


class OrganizerWorkbookError(ValueError):
    """The workbook does not satisfy the organizer dataset contract."""


@dataclass(frozen=True, slots=True)
class ParsedOrganizerDataset:
    """Validated organizer rows represented only in memory."""

    candidates: tuple[PositiveCandidateRecord, ...]
    annotations: tuple[OrganizerAnnotationRecord, ...]

    def __post_init__(self) -> None:
        if len(self.candidates) != EXPECTED_RECORD_COUNT:
            raise ValueError(f"expected {EXPECTED_RECORD_COUNT} candidates")
        if len(self.annotations) != EXPECTED_RECORD_COUNT:
            raise ValueError(f"expected {EXPECTED_RECORD_COUNT} annotations")
        candidate_ids = tuple(item.record_id for item in self.candidates)
        annotation_ids = tuple(item.record_id for item in self.annotations)
        if candidate_ids != annotation_ids:
            raise ValueError("candidate and annotation record IDs must match")


def _cell_name(row: int, column: int) -> str:
    return f"{chr(64 + column)}{row}"


def _required_text(value: object, row: int, column: int) -> str:
    cell = _cell_name(row, column)
    if not isinstance(value, str) or not value.strip():
        raise OrganizerWorkbookError(f"cell {cell} must contain text")
    return value


def _clean_identity_text(value: object, row: int, column: int) -> str:
    return " ".join(_required_text(value, row, column).split())


def _required_integer(value: object, row: int, column: int) -> int:
    cell = _cell_name(row, column)
    if isinstance(value, bool) or not isinstance(value, int):
        raise OrganizerWorkbookError(f"cell {cell} must contain an integer")
    return value


def _validate_layout(sheet: Worksheet) -> None:
    if sheet.max_row != DATA_END_ROW or sheet.max_column != LAST_COLUMN:
        raise OrganizerWorkbookError(
            f"expected data range B{HEADER_ROW}:J{DATA_END_ROW}, "
            f"found worksheet dimensions {sheet.max_row}x{sheet.max_column}"
        )

    actual_headers = tuple(
        sheet.cell(HEADER_ROW, column).value
        for column in range(FIRST_COLUMN, LAST_COLUMN + 1)
    )
    if actual_headers != EXPECTED_HEADERS:
        for offset, (expected, actual) in enumerate(
            zip(EXPECTED_HEADERS, actual_headers, strict=True)
        ):
            if actual != expected:
                column = FIRST_COLUMN + offset
                raise OrganizerWorkbookError(
                    f"unexpected header in {_cell_name(HEADER_ROW, column)}: "
                    f"expected {expected!r}, found {actual!r}"
                )


def _reject_formulas(sheet: Worksheet) -> None:
    for row in range(HEADER_ROW, DATA_END_ROW + 1):
        for column in range(FIRST_COLUMN, LAST_COLUMN + 1):
            cell = sheet.cell(row, column)
            if cell.data_type == "f":
                raise OrganizerWorkbookError(
                    f"formula is not allowed in {_cell_name(row, column)}"
                )


def _parse_row(
    sheet: Worksheet, row: int
) -> tuple[PositiveCandidateRecord, OrganizerAnnotationRecord]:
    source_number = _required_integer(sheet.cell(row, 2).value, row, 2)
    expected_number = row - DATA_START_ROW + 1
    if source_number != expected_number:
        raise OrganizerWorkbookError(
            f"cell {_cell_name(row, 2)} must contain sequence number {expected_number}"
        )

    canonical_name = _clean_identity_text(sheet.cell(row, 3).value, row, 3)
    domain = _clean_identity_text(sheet.cell(row, 4).value, row, 4)
    if domain not in ORGANIZER_SCOPE_KEYS:
        raise OrganizerWorkbookError(
            f"unsupported domain {domain!r} in {_cell_name(row, 4)}"
        )

    companies = _required_text(sheet.cell(row, 5).value, row, 5)
    rationale = _required_text(sheet.cell(row, 6).value, row, 6)
    stage = _required_text(sheet.cell(row, 7).value, row, 7)
    mention_trend = _required_text(sheet.cell(row, 8).value, row, 8)
    score = _required_integer(sheet.cell(row, 9).value, row, 9)
    if not 3 <= score <= 7:
        raise OrganizerWorkbookError(
            f"cell {_cell_name(row, 9)} must contain a score from 3 to 7"
        )
    sources = _required_text(sheet.cell(row, 10).value, row, 10)

    record_id = f"organizer-{source_number:03d}"
    candidate = PositiveCandidateRecord(
        record_id=record_id,
        source_row=row,
        canonical_name=canonical_name,
        domain=domain,
        analysis_scope_key=ORGANIZER_SCOPE_KEYS[domain],
        aliases=(),
        group_id=record_id,
        identity_status=IdentityStatus.PENDING_REVIEW,
        cutoff_date=ORGANIZER_CUTOFF_DATE,
    )
    annotation = OrganizerAnnotationRecord(
        record_id=record_id,
        source_number=source_number,
        companies_raw=companies,
        expert_rationale=rationale,
        expert_stage=stage,
        expert_mention_trend=mention_trend,
        expert_score=score,
        sources_raw=sources,
    )
    return candidate, annotation


def read_organizer_workbook(path: str | Path) -> ParsedOrganizerDataset:
    """Validate and parse the organizer XLSX without writing artifacts."""

    input_path = Path(path)
    if input_path.suffix.lower() != ".xlsx":
        raise OrganizerWorkbookError("organizer dataset must be an .xlsx file")
    if not input_path.is_file():
        raise OrganizerWorkbookError(f"organizer dataset not found: {input_path}")

    workbook = load_workbook(input_path, read_only=True, data_only=False)
    try:
        if SHEET_NAME not in workbook.sheetnames:
            raise OrganizerWorkbookError(f"worksheet {SHEET_NAME!r} was not found")
        sheet = workbook[SHEET_NAME]
        _validate_layout(sheet)
        _reject_formulas(sheet)

        candidates: list[PositiveCandidateRecord] = []
        annotations: list[OrganizerAnnotationRecord] = []
        for row in range(DATA_START_ROW, DATA_END_ROW + 1):
            candidate, annotation = _parse_row(sheet, row)
            candidates.append(candidate)
            annotations.append(annotation)

        return ParsedOrganizerDataset(tuple(candidates), tuple(annotations))
    finally:
        workbook.close()
